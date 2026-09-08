import streamlit as st
import os
import io
import re
import json
import zipfile
import shutil
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
import pytesseract
from PIL import Image
from docx import Document
import pandas as pd


# ============================================================
# STREAMLIT PAGE SETTINGS
# ============================================================

st.set_page_config(
    page_title="AI Resume Processor",
    page_icon="📄",
    layout="wide"
)


# ============================================================
# TESSERACT CONFIGURATION
# ============================================================

DEFAULT_TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

if os.path.exists(DEFAULT_TESSERACT_PATH):
    pytesseract.pytesseract.tesseract_cmd = DEFAULT_TESSERACT_PATH


# ============================================================
# BASIC CONFIGURATION
# ============================================================

SUPPORTED_EXTENSIONS = [".pdf", ".docx"]


# ============================================================
# SKILLS LIST
# ============================================================

SKILL_DATABASE = [
    # Programming
    "python",
    "java",
    "javascript",
    "typescript",
    "c",
    "c++",
    "c#",
    "php",
    "ruby",
    "go",
    "rust",
    "kotlin",
    "swift",

    # AI / ML
    "machine learning",
    "deep learning",
    "artificial intelligence",
    "computer vision",
    "natural language processing",
    "nlp",
    "generative ai",
    "genai",
    "llm",
    "large language models",
    "prompt engineering",
    "rag",

    # Libraries
    "numpy",
    "pandas",
    "matplotlib",
    "seaborn",
    "scikit-learn",
    "sklearn",
    "tensorflow",
    "keras",
    "pytorch",
    "opencv",
    "transformers",
    "hugging face",
    "langchain",

    # Web
    "html",
    "css",
    "react",
    "react.js",
    "angular",
    "vue",
    "node.js",
    "nodejs",
    "express",
    "django",
    "flask",
    "fastapi",
    "streamlit",

    # Database
    "sql",
    "mysql",
    "postgresql",
    "mongodb",
    "sqlite",
    "redis",
    "oracle",

    # Cloud / DevOps
    "aws",
    "azure",
    "gcp",
    "docker",
    "kubernetes",
    "jenkins",
    "github actions",
    "ci/cd",

    # Data
    "data science",
    "data analytics",
    "data analysis",
    "power bi",
    "tableau",
    "excel",
    "etl",
    "spark",
    "hadoop",

    # API / Backend
    "rest api",
    "restful api",
    "api development",
    "microservices",

    # Tools
    "git",
    "github",
    "jira",
    "linux",
    "postman"
]


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text):
    """
    Clean extracted resume text.
    """

    if not text:
        return ""

    # Remove null characters
    text = text.replace("\x00", " ")

    # Normalize line endings
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Remove too many spaces
    text = re.sub(r"[ \t]+", " ", text)

    # Remove repeated blank lines
    text = re.sub(r"\n\s*\n+", "\n\n", text)

    # Remove page marker lines
    text = re.sub(
        r"---\s*PAGE\s*\d+.*?---",
        "",
        text,
        flags=re.IGNORECASE
    )

    # Remove leading/trailing spaces
    lines = []

    for line in text.split("\n"):
        cleaned_line = line.strip()

        if cleaned_line:
            lines.append(cleaned_line)

    cleaned_text = "\n".join(lines)

    return cleaned_text.strip()


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(file_path):
    """
    Extract text from PDF.
    Uses normal PDF extraction first.
    Uses OCR if page has little/no text.
    """

    full_text = []

    try:
        pdf = fitz.open(file_path)

        for page_number, page in enumerate(pdf):

            text = page.get_text("text")

            # If normal extraction gives enough text
            if text and len(text.strip()) > 40:

                full_text.append(text)

            else:

                try:
                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(2, 2),
                        alpha=False
                    )

                    image_bytes = pix.tobytes("png")

                    image = Image.open(
                        io.BytesIO(image_bytes)
                    )

                    ocr_text = pytesseract.image_to_string(
                        image,
                        lang="eng"
                    )

                    full_text.append(ocr_text)

                except Exception as ocr_error:

                    full_text.append(
                        f"OCR_ERROR_PAGE_{page_number + 1}: "
                        f"{ocr_error}"
                    )

        pdf.close()

    except Exception as error:

        raise Exception(
            f"PDF extraction failed: {error}"
        )

    return "\n".join(full_text)


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(file_path):

    try:

        document = Document(file_path)

        text = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text.append(
                    paragraph.text.strip()
                )

        # Extract tables too
        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:

                    text.append(
                        " | ".join(row_text)
                    )

        return "\n".join(text)

    except Exception as error:

        raise Exception(
            f"DOCX extraction failed: {error}"
        )


# ============================================================
# EMAIL EXTRACTION
# ============================================================

def extract_email(text):

    email_pattern = r"""
    [A-Za-z0-9._%+-]+
    @
    [A-Za-z0-9.-]+
    \.
    [A-Za-z]{2,}
    """

    match = re.search(
        email_pattern,
        text,
        flags=re.VERBOSE
    )

    if match:
        return match.group(0)

    return None


# ============================================================
# PHONE EXTRACTION
# ============================================================

def extract_phone(text):

    phone_patterns = [
        r"\+91[\s\-]?[6-9]\d{9}",
        r"\+91[\s\-]?[6-9]\d{4}[\s\-]?\d{5}",
        r"\b[6-9]\d{9}\b",
        r"\b[6-9]\d{4}[\s\-]\d{5}\b"
    ]

    for pattern in phone_patterns:

        match = re.search(pattern, text)

        if match:

            phone = match.group(0)

            phone = re.sub(
                r"\s+",
                " ",
                phone
            )

            return phone

    return None


# ============================================================
# LINKEDIN EXTRACTION
# ============================================================

def extract_linkedin(text):

    pattern = r"""
    (?:
        https?://
    )?
    (?:
        www\.
    )?
    linkedin\.com/in/
    [A-Za-z0-9\-_%]+
    """

    match = re.search(
        pattern,
        text,
        flags=re.VERBOSE | re.IGNORECASE
    )

    if match:
        return match.group(0)

    return None


# ============================================================
# GITHUB EXTRACTION
# ============================================================

def extract_github(text):

    pattern = r"""
    (?:
        https?://
    )?
    (?:
        www\.
    )?
    github\.com/
    [A-Za-z0-9\-_.]+
    """

    match = re.search(
        pattern,
        text,
        flags=re.VERBOSE | re.IGNORECASE
    )

    if match:
        return match.group(0)

    return None


# ============================================================
# NAME EXTRACTION
# ============================================================

def extract_candidate_name(text):

    """
    Simple name extraction.

    We inspect first few lines and ignore lines
    containing contact/job related information.
    """

    lines = text.split("\n")

    ignored_words = [
        "resume",
        "curriculum",
        "email",
        "phone",
        "mobile",
        "linkedin",
        "github",
        "objective",
        "summary",
        "profile",
        "@",
        "http"
    ]

    for line in lines[:10]:

        candidate = line.strip()

        if not candidate:
            continue

        lower_candidate = candidate.lower()

        if any(
            word in lower_candidate
            for word in ignored_words
        ):
            continue

        if len(candidate) > 50:
            continue

        if len(candidate.split()) > 5:
            continue

        if re.search(r"\d", candidate):
            continue

        # Basic name validation
        if re.match(
            r"^[A-Za-z .'-]+$",
            candidate
        ):

            return candidate.title()

    return None


# ============================================================
# SKILLS EXTRACTION
# ============================================================

def extract_skills(text):

    detected_skills = []

    lower_text = text.lower()

    for skill in SKILL_DATABASE:

        pattern = r"\b" + re.escape(skill) + r"\b"

        if re.search(
            pattern,
            lower_text,
            re.IGNORECASE
        ):

            normalized_skill = skill.title()

            if normalized_skill not in detected_skills:

                detected_skills.append(
                    normalized_skill
                )

    return detected_skills


# ============================================================
# SECTION EXTRACTION
# ============================================================

SECTION_HEADERS = {
    "education": [
        "education",
        "academic background",
        "academic qualification",
        "qualifications"
    ],

    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "employment history",
        "internship",
        "internships"
    ],

    "projects": [
        "projects",
        "academic projects",
        "personal projects",
        "project experience"
    ],

    "certifications": [
        "certifications",
        "certificates",
        "courses",
        "training"
    ],

    "summary": [
        "summary",
        "professional summary",
        "profile summary",
        "career objective",
        "objective",
        "profile"
    ]
}


def get_section(text, section_name):

    lines = text.split("\n")

    headers = SECTION_HEADERS.get(
        section_name,
        []
    )

    all_headers = []

    for section_headers in SECTION_HEADERS.values():
        all_headers.extend(section_headers)

    section_lines = []

    capturing = False

    for line in lines:

        stripped_line = line.strip()

        lower_line = stripped_line.lower()

        normalized_line = re.sub(
            r"[^a-z ]",
            "",
            lower_line
        ).strip()

        if normalized_line in headers:

            capturing = True
            continue

        if capturing:

            if normalized_line in all_headers:

                break

            section_lines.append(
                stripped_line
            )

    section_text = "\n".join(
        section_lines
    ).strip()

    return section_text


# ============================================================
# EDUCATION EXTRACTION
# ============================================================

def extract_education(text):

    education_section = get_section(
        text,
        "education"
    )

    if not education_section:

        return []

    education_lines = education_section.split(
        "\n"
    )

    education = []

    degree_keywords = [
        "b.tech",
        "btech",
        "bachelor",
        "b.e",
        "be ",
        "b.sc",
        "bsc",
        "m.tech",
        "mtech",
        "master",
        "m.sc",
        "msc",
        "mba",
        "phd",
        "diploma",
        "12th",
        "10th"
    ]

    for line in education_lines:

        lower_line = line.lower()

        if any(
            keyword in lower_line
            for keyword in degree_keywords
        ):

            year_match = re.search(
                r"\b(19|20)\d{2}\b",
                line
            )

            education.append(
                {
                    "qualification": line,
                    "year": (
                        year_match.group(0)
                        if year_match
                        else None
                    )
                }
            )

    # If structured detection fails
    if not education and education_section:

        education.append(
            {
                "qualification":
                    education_section,
                "year": None
            }
        )

    return education


# ============================================================
# EXPERIENCE EXTRACTION
# ============================================================

def extract_experience(text):

    experience_section = get_section(
        text,
        "experience"
    )

    if not experience_section:

        return []

    lines = [
        line.strip()
        for line in experience_section.split("\n")
        if line.strip()
    ]

    experience = []

    for line in lines:

        if len(line) < 150:

            experience.append(
                {
                    "description": line
                }
            )

    return experience


# ============================================================
# PROJECT EXTRACTION
# ============================================================

def extract_projects(text):

    projects_section = get_section(
        text,
        "projects"
    )

    if not projects_section:

        return []

    lines = [
        line.strip()
        for line in projects_section.split("\n")
        if line.strip()
    ]

    projects = []

    for line in lines:

        projects.append(
            {
                "description": line
            }
        )

    return projects


# ============================================================
# CERTIFICATION EXTRACTION
# ============================================================

def extract_certifications(text):

    certification_section = get_section(
        text,
        "certifications"
    )

    if not certification_section:

        return []

    lines = [
        line.strip()
        for line
        in certification_section.split("\n")
        if line.strip()
    ]

    return lines


# ============================================================
# EXPERIENCE YEARS EXTRACTION
# ============================================================

def extract_total_experience(text):

    patterns = [
        r"(\d+(?:\.\d+)?)\+?\s*years?\s+of\s+experience",
        r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?)\+?\s*years?",
        r"(\d+(?:\.\d+)?)\+?\s*yrs?\s+experience"
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            re.IGNORECASE
        )

        if match:

            try:
                return float(
                    match.group(1)
                )

            except:
                pass

    return None


# ============================================================
# RESUME → STRUCTURED JSON
# ============================================================

def convert_resume_to_json(
    text,
    filename,
    resume_id
):

    candidate_data = {

        "resume_id": resume_id,

        "source_file": filename,

        "candidate_name":
            extract_candidate_name(text),

        "contact": {
            "email":
                extract_email(text),

            "phone":
                extract_phone(text),

            "linkedin":
                extract_linkedin(text),

            "github":
                extract_github(text)
        },

        "skills":
            extract_skills(text),

        "professional_summary":
            get_section(
                text,
                "summary"
            ),

        "education":
            extract_education(text),

        "experience":
            extract_experience(text),

        "total_experience_years":
            extract_total_experience(text),

        "projects":
            extract_projects(text),

        "certifications":
            extract_certifications(text),

        "raw_text":
            text
    }

    return candidate_data


# ============================================================
# GET RESUME FILES FROM ZIP
# ============================================================

def extract_zip(uploaded_zip):

    temp_directory = tempfile.mkdtemp()

    zip_path = os.path.join(
        temp_directory,
        "uploaded_resumes.zip"
    )

    with open(
        zip_path,
        "wb"
    ) as file:

        file.write(
            uploaded_zip.getbuffer()
        )

    extraction_folder = os.path.join(
        temp_directory,
        "resumes"
    )

    os.makedirs(
        extraction_folder,
        exist_ok=True
    )

    with zipfile.ZipFile(
        zip_path,
        "r"
    ) as zip_ref:

        zip_ref.extractall(
            extraction_folder
        )

    resume_files = []

    for root, directories, files in os.walk(
        extraction_folder
    ):

        for filename in files:

            file_path = os.path.join(
                root,
                filename
            )

            extension = Path(
                filename
            ).suffix.lower()

            if extension in SUPPORTED_EXTENSIONS:

                resume_files.append(
                    file_path
                )

    return (
        temp_directory,
        resume_files
    )


# ============================================================
# PROCESS SINGLE RESUME
# ============================================================

def process_resume(
    file_path,
    index
):

    filename = os.path.basename(
        file_path
    )

    extension = Path(
        filename
    ).suffix.lower()

    if extension == ".pdf":

        raw_text = extract_text_from_pdf(
            file_path
        )

    elif extension == ".docx":

        raw_text = extract_text_from_docx(
            file_path
        )

    else:

        raise ValueError(
            "Unsupported file format"
        )

    cleaned_text = clean_text(
        raw_text
    )

    resume_id = (
        f"resume_{index:03d}"
    )

    structured_data = (
        convert_resume_to_json(
            cleaned_text,
            filename,
            resume_id
        )
    )

    return structured_data


# ============================================================
# JSON ZIP CREATION
# ============================================================

def create_json_zip(
    processed_resumes
):

    memory_zip = io.BytesIO()

    with zipfile.ZipFile(
        memory_zip,
        "w",
        zipfile.ZIP_DEFLATED
    ) as zip_file:

        for resume in processed_resumes:

            filename = (
                f"{resume['resume_id']}.json"
            )

            json_data = json.dumps(
                resume,
                indent=4,
                ensure_ascii=False
            )

            zip_file.writestr(
                filename,
                json_data
            )

    memory_zip.seek(0)

    return memory_zip


# ============================================================
# STREAMLIT UI
# ============================================================

st.title(
    "📄 AI Resume Processing System"
)

st.write(
    """
    Upload a ZIP file containing PDF/DOCX resumes.

    The system will automatically:

    **Extract → OCR → Clean → Structure → Convert to JSON**
    """
)

st.divider()


# ============================================================
# TESSERACT STATUS
# ============================================================

if os.path.exists(
    DEFAULT_TESSERACT_PATH
):

    st.success(
        "✅ Tesseract OCR detected"
    )

else:

    st.warning(
        """
        ⚠️ Tesseract was not found at:

        C:\\Program Files\\Tesseract-OCR\\tesseract.exe

        Normal PDFs and DOCX files will still work,
        but scanned PDFs may fail OCR.
        """
    )


# ============================================================
# FILE UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "Upload Resume ZIP",
    type=["zip"],
    help=(
        "Upload one ZIP file containing "
        "PDF and/or DOCX resumes."
    )
)


# ============================================================
# PROCESS BUTTON
# ============================================================

if uploaded_file is not None:

    st.info(
        f"Uploaded file: {uploaded_file.name}"
    )

    if st.button(
        "🚀 Process Resumes",
        type="primary",
        use_container_width=True
    ):

        temp_directory = None

        try:

            with st.spinner(
                "Extracting ZIP file..."
            ):

                (
                    temp_directory,
                    resume_files
                ) = extract_zip(
                    uploaded_file
                )

            total_resumes = len(
                resume_files
            )

            if total_resumes == 0:

                st.error(
                    """
                    No PDF or DOCX resumes were found
                    inside the ZIP file.
                    """
                )

                st.stop()

            st.success(
                f"✅ Found {total_resumes} resumes"
            )

            processed_resumes = []

            failed_resumes = []

            progress_bar = st.progress(
                0
            )

            status_text = st.empty()

            # =================================================
            # PROCESS ALL RESUMES
            # =================================================

            for index, file_path in enumerate(
                resume_files,
                start=1
            ):

                filename = os.path.basename(
                    file_path
                )

                status_text.write(
                    f"Processing {index}/{total_resumes}: "
                    f"**{filename}**"
                )

                try:

                    structured_resume = (
                        process_resume(
                            file_path,
                            index
                        )
                    )

                    processed_resumes.append(
                        structured_resume
                    )

                except Exception as error:

                    failed_resumes.append(
                        {
                            "filename":
                                filename,

                            "error":
                                str(error)
                        }
                    )

                progress = (
                    index /
                    total_resumes
                )

                progress_bar.progress(
                    progress
                )

            status_text.write(
                "✅ Processing completed"
            )

            # Store in Streamlit session
            st.session_state[
                "processed_resumes"
            ] = processed_resumes

            st.session_state[
                "failed_resumes"
            ] = failed_resumes

            # =================================================
            # SUMMARY
            # =================================================

            st.divider()

            st.subheader(
                "📊 Processing Summary"
            )

            col1, col2, col3 = st.columns(
                3
            )

            col1.metric(
                "Total Resumes",
                total_resumes
            )

            col2.metric(
                "Successfully Processed",
                len(processed_resumes)
            )

            col3.metric(
                "Failed",
                len(failed_resumes)
            )

            if failed_resumes:

                st.warning(
                    "Some resumes failed to process."
                )

                failed_df = pd.DataFrame(
                    failed_resumes
                )

                st.dataframe(
                    failed_df,
                    use_container_width=True
                )

        finally:

            if temp_directory:

                try:

                    shutil.rmtree(
                        temp_directory
                    )

                except:
                    pass


# ============================================================
# DISPLAY RESULTS
# ============================================================

if (
    "processed_resumes"
    in st.session_state
):

    processed_resumes = (
        st.session_state[
            "processed_resumes"
        ]
    )

    if processed_resumes:

        st.divider()

        st.header(
            "👥 Extracted Candidates"
        )

        # ====================================================
        # CANDIDATE TABLE
        # ====================================================

        table_data = []

        for resume in processed_resumes:

            table_data.append(
                {
                    "Resume ID":
                        resume[
                            "resume_id"
                        ],

                    "Name":
                        resume[
                            "candidate_name"
                        ],

                    "Email":
                        resume[
                            "contact"
                        ][
                            "email"
                        ],

                    "Phone":
                        resume[
                            "contact"
                        ][
                            "phone"
                        ],

                    "Skills Count":
                        len(
                            resume[
                                "skills"
                            ]
                        ),

                    "Experience":
                        resume[
                            "total_experience_years"
                        ]
                }
            )

        candidate_df = pd.DataFrame(
            table_data
        )

        st.dataframe(
            candidate_df,
            use_container_width=True
        )

        # ====================================================
        # SELECT CANDIDATE
        # ====================================================

        st.subheader(
            "🔍 Candidate Details"
        )

        resume_options = {
            (
                f"{resume['resume_id']} - "
                f"{resume['candidate_name'] or resume['source_file']}"
            ):
            resume

            for resume
            in processed_resumes
        }

        selected_candidate = (
            st.selectbox(
                "Select Candidate",
                list(
                    resume_options.keys()
                )
            )
        )

        candidate = (
            resume_options[
                selected_candidate
            ]
        )

        # ====================================================
        # BASIC DETAILS
        # ====================================================

        col1, col2 = st.columns(
            2
        )

        with col1:

            st.write(
                "### Personal Information"
            )

            st.write(
                "**Name:**",
                candidate[
                    "candidate_name"
                ]
                or
                "Not detected"
            )

            st.write(
                "**Email:**",
                candidate[
                    "contact"
                ][
                    "email"
                ]
                or
                "Not detected"
            )

            st.write(
                "**Phone:**",
                candidate[
                    "contact"
                ][
                    "phone"
                ]
                or
                "Not detected"
            )

        with col2:

            st.write(
                "### Links"
            )

            st.write(
                "**LinkedIn:**",
                candidate[
                    "contact"
                ][
                    "linkedin"
                ]
                or
                "Not detected"
            )

            st.write(
                "**GitHub:**",
                candidate[
                    "contact"
                ][
                    "github"
                ]
                or
                "Not detected"
            )

            st.write(
                "**Experience:**",
                candidate[
                    "total_experience_years"
                ]
                or
                "Not detected"
            )

        # ====================================================
        # SKILLS
        # ====================================================

        st.write(
            "### 🛠 Skills"
        )

        if candidate[
            "skills"
        ]:

            st.write(
                ", ".join(
                    candidate[
                        "skills"
                    ]
                )
            )

        else:

            st.write(
                "No skills detected."
            )

        # ====================================================
        # EDUCATION
        # ====================================================

        st.write(
            "### 🎓 Education"
        )

        if candidate[
            "education"
        ]:

            for education in candidate[
                "education"
            ]:

                st.write(
                    "•",
                    education[
                        "qualification"
                    ]
                )

        else:

            st.write(
                "No education information detected."
            )

        # ====================================================
        # EXPERIENCE
        # ====================================================

        st.write(
            "### 💼 Experience"
        )

        if candidate[
            "experience"
        ]:

            for experience in candidate[
                "experience"
            ]:

                st.write(
                    "•",
                    experience[
                        "description"
                    ]
                )

        else:

            st.write(
                "No experience information detected."
            )

        # ====================================================
        # PROJECTS
        # ====================================================

        st.write(
            "### 📁 Projects"
        )

        if candidate[
            "projects"
        ]:

            for project in candidate[
                "projects"
            ]:

                st.write(
                    "•",
                    project[
                        "description"
                    ]
                )

        else:

            st.write(
                "No project information detected."
            )

        # ====================================================
        # STRUCTURED JSON
        # ====================================================

        with st.expander(
            "🧾 View Structured JSON"
        ):

            st.json(
                candidate
            )

        # ====================================================
        # RAW TEXT
        # ====================================================

        with st.expander(
            "📄 View Cleaned Resume Text"
        ):

            st.text_area(
                "Extracted Text",
                candidate[
                    "raw_text"
                ],
                height=400
            )

        # ====================================================
        # DOWNLOAD SINGLE JSON
        # ====================================================

        candidate_json = json.dumps(
            candidate,
            indent=4,
            ensure_ascii=False
        )

        st.download_button(
            label="⬇️ Download Selected Candidate JSON",
            data=candidate_json,
            file_name=(
                f"{candidate['resume_id']}.json"
            ),
            mime="application/json"
        )

        # ====================================================
        # DOWNLOAD ALL JSON FILES
        # ====================================================

        all_json_zip = create_json_zip(
            processed_resumes
        )

        st.download_button(
            label="📦 Download All Resume JSON Files",
            data=all_json_zip,
            file_name="processed_resumes_json.zip",
            mime="application/zip",
            use_container_width=True
        )

        # ====================================================
        # DOWNLOAD COMBINED JSON
        # ====================================================

        combined_json = json.dumps(
            processed_resumes,
            indent=4,
            ensure_ascii=False
        )

        st.download_button(
            label="⬇️ Download All Candidates as One JSON",
            data=combined_json,
            file_name="all_candidates.json",
            mime="application/json",
            use_container_width=True
        )